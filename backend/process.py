import os
import json
import re
import time
from groq import Groq
from PyPDF2 import PdfReader
from pdf2image import convert_from_path
import pytesseract
from dotenv import load_dotenv

load_dotenv()

# ============================================================
# CONFIGURATION
# ============================================================

CHUNK_SIZE = 5000
CHUNK_OVERLAP = 500

GROQ_MODEL = "qwen/qwen3.6-27b"
GROQ_RETRIES = 3
GROQ_RETRY_DELAY = 2.0
MIN_SECONDS_BETWEEN_GROQ_CALLS = 1.0

ALLOWED_CATEGORIES = [
    "Tender / Bid Document",
    "Maintenance Log / Work Order",
    "Inspection Report",
    "Incident Report",
    "Policy / Standard Operating Procedure",
    "Financial / Budget Document",
    "HR / Personnel Record",
    "Other",
]

_last_groq_call_time = 0.0

# ============================================================
# GROQ CLIENT
# ============================================================

def get_groq_client():
    api_key = os.getenv("GROQ_API_KEY")
    if not api_key:
        return None
    return Groq(api_key=api_key)

def _wait_before_groq_call():
    global _last_groq_call_time
    now = time.time()
    elapsed = now - _last_groq_call_time
    if elapsed < MIN_SECONDS_BETWEEN_GROQ_CALLS:
        time.sleep(MIN_SECONDS_BETWEEN_GROQ_CALLS - elapsed)
    _last_groq_call_time = time.time()

# ============================================================
# PAGE COUNT
# ============================================================

def get_page_count(pdf_path):
    try:
        reader = PdfReader(pdf_path)
        return len(reader.pages)
    except Exception as e:
        print(f"⚠️ PyPDF2 page count failed: {e}")
        try:
            images = convert_from_path(pdf_path, dpi=100)
            return len(images)
        except Exception as fallback_error:
            print(f"❌ PDF page count failed: {fallback_error}")
            return 0

# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_native_pdf_text(pdf_path):
    pages = []
    try:
        reader = PdfReader(pdf_path)
        for page_number, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
                pages.append({
                    "page": page_number,
                    "text": text.strip(),
                })
            except Exception as e:
                print(f"⚠️ Could not extract page {page_number}: {e}")
                pages.append({
                    "page": page_number,
                    "text": "",
                })
    except Exception as e:
        print(f"❌ Native PDF extraction failed: {e}")
        return []
    return pages

def ocr_pdf_page(pdf_path, page_number, dpi=200):
    try:
        images = convert_from_path(
            pdf_path,
            dpi=dpi,
            first_page=page_number,
            last_page=page_number,
        )
        if not images:
            return ""
        return pytesseract.image_to_string(images[0]).strip()
    except Exception as e:
        print(f"❌ OCR failed on page {page_number}: {e}")
        return ""

def extract_text_from_pdf(pdf_path):
    print(f"📄 Extracting PDF: {pdf_path}")
    native_pages = extract_native_pdf_text(pdf_path)

    if not native_pages:
        page_count = get_page_count(pdf_path)
        native_pages = [
            {"page": i, "text": ""}
            for i in range(1, page_count + 1)
        ]

    extracted_pages = []

    for page_info in native_pages:
        page_number = page_info["page"]
        text = page_info["text"]

        if len(text.strip()) < 50:
            print(f"🔎 Page {page_number}: little/no native text -> OCR")
            ocr_text = ocr_pdf_page(pdf_path, page_number)
            if ocr_text:
                text = ocr_text
                print(f"✅ OCR successful on page {page_number}")
            else:
                print(f"⚠️ OCR produced no text on page {page_number}")
        else:
            print(f"✅ Page {page_number}: native text extracted")

        extracted_pages.append((page_number, text))

    parts = []

    for page_number, text in extracted_pages:
        parts.append(f"\n--- Page {page_number} ---\n{text}")

    return "\n".join(parts).strip()

# ============================================================
# IMAGE OCR
# ============================================================

def extract_text_from_image(image_path):
    from PIL import Image
    try:
        with Image.open(image_path) as image:
            return pytesseract.image_to_string(image).strip()
    except Exception as e:
        print(f"❌ Image OCR error: {e}")
        return ""

# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_text_from_docx(docx_path):
    try:
        from docx import Document
        doc = Document(docx_path)
        parts = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        for table_index, table in enumerate(doc.tables, start=1):
            parts.append(f"\n--- Table {table_index} ---")
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cell_text = cell.text.strip().replace("\n", " ")
                    cells.append(cell_text)
                parts.append(" | ".join(cells))

        return "\n".join(parts) if parts else ""
    except Exception as e:
        print(f"❌ DOCX extraction error: {e}")
        return ""

# ============================================================
# EVIDENCE / MATH SCORER
# ============================================================

def calculate_evidence_score(text):
    category_keywords = {
        "Tender / Bid Document": [
            "tender", "bid", "procurement", "submission",
            "contract", "rfp", "quotation", "bidding",
        ],
        "Maintenance Log / Work Order": [
            "maintenance", "repair", "work order",
            "inspection", "track", "equipment", "service",
        ],
        "Inspection Report": [
            "inspection", "safety", "audit", "compliance",
            "check", "report", "condition",
        ],
        "Incident Report": [
            "incident", "accident", "near-miss",
            "emergency", "injury", "damage",
        ],
        "Policy / Standard Operating Procedure": [
            "policy", "procedure", "sop", "guideline",
            "protocol", "standard",
        ],
        "Financial / Budget Document": [
            "budget", "financial", "cost", "expenditure",
            "account", "fund", "estimate",
        ],
        "HR / Personnel Record": [
            "employee", "hr", "personnel", "leave",
            "attendance", "recruitment", "staff",
        ],
        "Other": [
            "form", "application", "miscellaneous", "unknown",
        ],
    }

    text_lower = text.lower()
    scores = {}

    for category, keywords in category_keywords.items():
        score = 0.40
        matched_keywords = 0

        for keyword in keywords:
            if keyword.lower() in text_lower:
                matched_keywords += 1
                score += 0.10

        scores[category] = min(score, 1.0)

    return scores

# ============================================================
# CHUNKING
# ============================================================

def _split_long_text(text, chunk_size, overlap):
    pieces = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))

        if end < len(text):
            candidates = [
                text.rfind("\n\n", start, end),
                text.rfind("\n", start, end),
                text.rfind(". ", start, end),
                text.rfind(" ", start, end),
            ]

            best = max(candidates)

            if best > start + int(chunk_size * 0.65):
                end = best + (
                    2 if text[best:best + 2] == "\n\n" else 1
                )

        pieces.append(text[start:end].strip())

        if end >= len(text):
            break

        start = max(end - overlap, start + 1)

    return [p for p in pieces if p]

def create_chunks(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    if not text or not text.strip():
        return []

    text = text.strip()

    if len(text) <= chunk_size:
        return [{
            "chunk_number": 1,
            "total_chunks": 1,
            "text": text,
        }]

    page_pattern = r"(?=--- Page \d+ ---)"
    pages = [
        p.strip()
        for p in re.split(page_pattern, text)
        if p.strip()
    ]

    raw_chunks = []
    current = ""

    for page in pages:
        if len(page) > chunk_size:
            if current:
                raw_chunks.append(current)
                current = ""

            raw_chunks.extend(
                _split_long_text(
                    page,
                    chunk_size,
                    overlap,
                )
            )
            continue

        candidate = (
            f"{current}\n\n{page}"
            if current
            else page
        )

        if len(candidate) <= chunk_size:
            current = candidate
        else:
            if current:
                raw_chunks.append(current)
            current = page

    if current:
        raw_chunks.append(current)

    final_chunks = []

    for index, chunk in enumerate(raw_chunks):
        if index == 0:
            final_chunks.append(chunk)
            continue

        previous_tail = raw_chunks[index - 1][-overlap:]

        combined = (
            f"[Previous context]\n"
            f"{previous_tail}\n\n"
            f"[Current chunk]\n"
            f"{chunk}"
        )

        if len(combined) <= chunk_size + overlap + 100:
            final_chunks.append(combined)
        else:
            final_chunks.append(chunk)

    total = len(final_chunks)

    return [
        {
            "chunk_number": index + 1,
            "total_chunks": total,
            "text": chunk,
        }
        for index, chunk in enumerate(final_chunks)
    ]

# ============================================================
# JSON / NORMALIZATION
# ============================================================

def clean_json_response(response_text):
    if not response_text:
        raise ValueError("Empty Groq response")

    text = response_text.strip()

    text = re.sub(
        r"^```(?:json)?\s*",
        "",
        text,
        flags=re.IGNORECASE,
    )

    text = re.sub(
        r"\s*```$",
        "",
        text,
    )

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(
            r"\{.*\}",
            text,
            re.DOTALL,
        )

        if not match:
            raise

        return json.loads(match.group())

def normalize_action_items(action_items):
    if not isinstance(action_items, list):
        return []

    normalized = []

    for item in action_items:
        if isinstance(item, str):
            value = item.strip()
        elif isinstance(item, dict):
            value = (
                item.get("task")
                or item.get("action")
                or item.get("description")
                or item.get("item")
                or ""
            )
            value = str(value).strip()
        else:
            value = ""

        if value:
            normalized.append(value)

    return normalized

def deduplicate_action_items(action_items):
    result = []
    seen = set()

    for item in action_items:
        cleaned = re.sub(
            r"\s+",
            " ",
            str(item).strip(),
        )

        if not cleaned:
            continue

        normalized = re.sub(
            r"[^a-z0-9 ]",
            "",
            cleaned.lower(),
        )

        if not normalized or normalized in seen:
            continue

        words = set(normalized.split())
        duplicate = False

        for existing in result:
            existing_normalized = re.sub(
                r"[^a-z0-9 ]",
                "",
                existing.lower(),
            )

            existing_words = set(
                existing_normalized.split()
            )

            if not existing_words:
                continue

            intersection = len(words & existing_words)
            smaller = min(
                len(words),
                len(existing_words),
            )

            if smaller and intersection / smaller >= 0.85:
                duplicate = True
                break

        if not duplicate:
            result.append(cleaned)
            seen.add(normalized)

    return result

def normalize_deadlines(value):
    if value is None:
        return []

    values = value if isinstance(value, list) else [value]
    result = []

    for item in values:
        value = str(item).strip()

        if value and value.lower() != "no deadline specified":
            if value not in result:
                result.append(value)

    return result

def normalize_chunk_result(result):
    if not isinstance(result, dict):
        raise ValueError(
            "Groq returned a non-object JSON response"
        )

    category = result.get("category", "Other")

    if category not in ALLOWED_CATEGORIES:
        category = "Other"

    summary = str(
        result.get("summary", "")
    ).strip()

    action_items = normalize_action_items(
        result.get("action_items", [])
    )

    deadlines = normalize_deadlines(
        result.get(
            "deadlines",
            result.get("deadline"),
        )
    )

    try:
        confidence = float(
            result.get("confidence", 0.5)
        )
    except (TypeError, ValueError):
        confidence = 0.5

    confidence = max(
        0.0,
        min(1.0, confidence),
    )

    return {
        "category": category,
        "summary": summary,
        "action_items": action_items,
        "deadlines": deadlines,
        "confidence": confidence,
    }

# ============================================================
# GROQ JSON REQUEST
# ============================================================

def _groq_request_json(client, prompt):
    errors = []

    try:
        _wait_before_groq_call()
        print("🤖 Groq JSON mode")

        completion = client.chat.completions.create(
            model=GROQ_MODEL,
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are a KMRL document extraction engine. "
                        "Return ONLY one valid JSON object. "
                        "Do not return markdown. "
                        "Do not return explanations."
                    ),
                },
                {
                    "role": "user",
                    "content": prompt,
                },
            ],
            temperature=0.1,
            max_completion_tokens=1200,
            response_format={
                "type": "json_object",
            },
            reasoning_format="hidden",
            reasoning_effort="none",
            stream=False,
        )

        response_text = (
            completion.choices[0].message.content
            or ""
        )

        if not response_text.strip():
            raise ValueError(
                "Groq returned an empty JSON response"
            )

        result = clean_json_response(response_text)

        print("✅ Groq JSON mode succeeded")
        return result

    except Exception as e:
        errors.append(f"JSON mode: {str(e)}")
        print(f"⚠️ Groq JSON mode failed: {e}")

    try:
        _wait_before_groq_call()
        print("🤖 Groq plain JSON mode")

        completion = client.chat.completions.create(
            model=GROQ_MODEL,
            messages=[
                {
                    "role": "user",
                    "content": (
                        "Return ONLY valid JSON. "
                        "Do not use markdown. "
                        "Do not explain anything.\n\n"
                        + prompt
                    ),
                },
            ],
            temperature=0.1,
            max_completion_tokens=1200,
            reasoning_format="hidden",
            reasoning_effort="none",
            stream=False,
        )

        response_text = (
            completion.choices[0].message.content
            or ""
        )

        if not response_text.strip():
            raise ValueError(
                "Groq plain mode returned an empty response"
            )

        result = clean_json_response(response_text)

        print("✅ Groq plain JSON mode succeeded")
        return result

    except Exception as e:
        errors.append(f"Plain mode: {str(e)}")
        print(f"⚠️ Groq plain mode failed: {e}")

    try:
        _wait_before_groq_call()
        print("🤖 Groq simplified extraction attempt")

        simple_prompt = f"""
Analyze this KMRL document chunk.

Return ONLY this JSON object:

{{
  "category": "Other",
  "summary": "",
  "action_items": [],
  "deadlines": [],
  "confidence": 0.0
}}

Allowed categories:

Tender / Bid Document
Maintenance Log / Work Order
Inspection Report
Incident Report
Policy / Standard Operating Procedure
Financial / Budget Document
HR / Personnel Record
Other

Rules:

- Use only information in the document.
- Never invent information.
- Extract actual action items.
- Extract explicit deadlines.
- Use an empty array if none exist.
- Return ONLY JSON.

DOCUMENT:

{prompt}
"""

        completion = client.chat.completions.create(
            model=GROQ_MODEL,
            messages=[
                {
                    "role": "user",
                    "content": simple_prompt,
                },
            ],
            temperature=0.0,
            max_completion_tokens=1000,
            response_format={
                "type": "json_object",
            },
            reasoning_format="hidden",
            reasoning_effort="none",
            stream=False,
        )

        response_text = (
            completion.choices[0].message.content
            or ""
        )

        if not response_text.strip():
            raise ValueError(
                "Simplified Groq response was empty"
            )

        result = clean_json_response(response_text)

        print("✅ Groq simplified extraction succeeded")
        return result

    except Exception as e:
        errors.append(f"Simplified mode: {str(e)}")
        print(f"⚠️ Groq simplified mode failed: {e}")

    combined_error = " | ".join(errors)

    print("❌ All Groq request methods failed.")
    print(f"   {combined_error}")

    raise RuntimeError(combined_error)

# ============================================================
# SINGLE CHUNK ANALYSIS
# ============================================================

def analyze_chunk_with_groq(
    chunk_text,
    chunk_number,
    total_chunks,
):
    client = get_groq_client()

    if client is None:
        print("⚠️ GROQ_API_KEY is missing.")

        return {
            "success": False,
            "category": "Other",
            "summary": "",
            "action_items": [],
            "deadlines": [],
            "confidence": 0.0,
            "error": "GROQ_API_KEY is missing",
        }

    prompt = f"""
You are analyzing one part of an official
Kochi Metro Rail Limited (KMRL) document.

This is chunk {chunk_number} of {total_chunks}.

IMPORTANT:

Analyze ONLY the supplied document text.

Never invent:

- names
- dates
- deadlines
- amounts
- departments
- locations
- reference numbers
- tasks
- facts

OCR text may contain small errors.
Correct an obvious OCR error only when the surrounding
text makes the intended meaning completely clear.

CLASSIFICATION:

Choose exactly ONE category:

- Tender / Bid Document
- Maintenance Log / Work Order
- Inspection Report
- Incident Report
- Policy / Standard Operating Procedure
- Financial / Budget Document
- HR / Personnel Record
- Other

ACTION ITEMS:

Extract actual required actions.

Examples:

- submit a document
- complete a repair
- perform an inspection
- obtain approval
- provide information
- review something
- prepare something
- make a payment
- complete assigned work

Do NOT turn a historical event into an action item.

Do NOT create an action item merely because a word such
as "should" appears in a descriptive sentence.

DEADLINES:

Extract explicit dates that represent:

- due dates
- submission dates
- completion dates
- expiry dates
- required action dates

Preserve the date as written.

If there are no deadlines:

"deadlines": []

SUMMARY:

Write a short factual summary of this chunk.

CONFIDENCE:

Use a value between 0.0 and 1.0.

Return ONLY one JSON object.

Required structure:

{{
    "category": "one allowed category",
    "summary": "short factual summary",
    "action_items": [],
    "deadlines": [],
    "confidence": 0.0
}}

DOCUMENT CHUNK:

{chunk_text}
"""

    last_error = None

    for attempt in range(1, GROQ_RETRIES + 1):
        try:
            print(
                f"🤖 Groq chunk "
                f"{chunk_number}/{total_chunks} "
                f"(attempt {attempt})"
            )

            result = _groq_request_json(
                client,
                prompt,
            )

            normalized = normalize_chunk_result(result)
            normalized["success"] = True

            print(
                f"✅ Groq chunk "
                f"{chunk_number}/{total_chunks} succeeded"
            )

            return normalized

        except Exception as e:
            last_error = str(e)

            print(
                f"⚠️ Chunk "
                f"{chunk_number}/{total_chunks} failed: "
                f"{e}"
            )

            if attempt < GROQ_RETRIES:
                delay = GROQ_RETRY_DELAY * attempt

                print(
                    f"⏳ Retrying in "
                    f"{delay} seconds..."
                )

                time.sleep(delay)

    print(
        f"🛡️ Chunk {chunk_number} "
        f"will use local fallback."
    )

    return {
        "success": False,
        "category": "Other",
        "summary": "",
        "action_items": [],
        "deadlines": [],
        "confidence": 0.0,
        "error": last_error,
    }

# ============================================================
# LOCAL FALLBACK EXTRACTION
# ============================================================

def extract_deadlines_locally(text):
    if not text:
        return []

    patterns = [
        r"\b\d{4}[-/]\d{1,2}[-/]\d{1,2}\b",
        r"\b\d{1,2}[-/]\d{1,2}[-/]\d{4}\b",
        r"\b\d{1,2}\s+(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{4}\b",
        r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+\d{4}\b",
    ]

    found = []

    for pattern in patterns:
        for match in re.findall(
            pattern,
            text,
            flags=re.IGNORECASE,
        ):
            value = match.strip()

            if value not in found:
                found.append(value)

    return found[:20]

def extract_action_candidates_locally(text):
    if not text:
        return []

    candidates = []

    action_pattern = re.compile(
        r"(?im)^\s*(?:"
        r"action(?: required)?|"
        r"required action|"
        r"to be (?:done|completed|submitted|carried out)|"
        r"shall|"
        r"must|"
        r"should|"
        r"submit|"
        r"complete|"
        r"repair|"
        r"inspect|"
        r"approve|"
        r"review|"
        r"provide|"
        r"prepare"
        r")\b.*$"
    )

    for line in text.splitlines():
        line = re.sub(
            r"^\s*[-*•\d.)]+\s*",
            "",
            line.strip(),
        )

        if not line:
            continue

        if action_pattern.match(line):
            if 10 <= len(line) <= 400:
                candidates.append(line)

    return deduplicate_action_items(candidates)[:20]

# ============================================================
# SYNTHESIS
# ============================================================

def synthesize_results(chunk_results, full_text):
    successful = [
        r
        for r in chunk_results
        if r.get("success")
    ]

    evidence_scores = calculate_evidence_score(full_text)

    best_evidence_category = max(
        evidence_scores,
        key=evidence_scores.get,
    )

    if not successful:
        deadlines = extract_deadlines_locally(full_text)
        actions = extract_action_candidates_locally(full_text)

        return {
            "category": best_evidence_category,
            "summary": (
                "The document was processed using the "
                "local document-evidence fallback because "
                "the AI service was temporarily unavailable."
            ),
            "action_items": actions,
            "deadline": (
                "; ".join(deadlines)
                if deadlines
                else "No deadline specified"
            ),
            "groq_confidence": 0.0,
            "successful_chunks": 0,
        }

    category_scores = {}

    for result in successful:
        category = result.get(
            "category",
            "Other",
        )

        confidence = float(
            result.get(
                "confidence",
                0.0,
            )
        )

        category_scores[category] = (
            category_scores.get(
                category,
                0.0,
            )
            + confidence
        )

    groq_category = (
        max(
            category_scores,
            key=category_scores.get,
        )
        if category_scores
        else "Other"
    )

    summaries = [
        r.get("summary", "").strip()
        for r in successful
        if r.get("summary")
    ]

    all_actions = []

    for result in successful:
        all_actions.extend(
            result.get(
                "action_items",
                [],
            )
        )

    all_actions = deduplicate_action_items(
        all_actions
    )

    all_deadlines = []

    for result in successful:
        all_deadlines.extend(
            result.get(
                "deadlines",
                [],
            )
        )

    all_deadlines = [
        d
        for d in all_deadlines
        if d
    ]

    if not all_deadlines:
        all_deadlines = extract_deadlines_locally(
            full_text
        )

    unique_deadlines = []

    for deadline in all_deadlines:
        cleaned = re.sub(
            r"\s+",
            " ",
            str(deadline).strip(),
        )

        if (
            cleaned
            and cleaned not in unique_deadlines
        ):
            unique_deadlines.append(cleaned)

    deadline = (
        "; ".join(unique_deadlines)
        if unique_deadlines
        else "No deadline specified"
    )

    if not all_actions:
        all_actions = extract_action_candidates_locally(
            full_text
        )

    confidences = [
        float(
            r.get(
                "confidence",
                0.0,
            )
        )
        for r in successful
    ]

    groq_confidence = (
        sum(confidences) / len(confidences)
        if confidences
        else 0.0
    )

    summary_parts = []

    for summary in summaries:
        if summary and summary not in summary_parts:
            summary_parts.append(summary)

        if len(summary_parts) >= 3:
            break

    if summary_parts:
        final_summary = " ".join(summary_parts)
    else:
        final_summary = (
            "The document was analyzed from its "
            "available extracted text."
        )

    return {
        "category": groq_category,
        "summary": final_summary,
        "action_items": all_actions,
        "deadline": deadline,
        "groq_confidence": groq_confidence,
        "successful_chunks": len(successful),
    }

# ============================================================
# MAIN HYBRID ANALYSIS
# ============================================================

def analyze_with_groq(text):
    if not text or not text.strip():
        return {
            "category": "Other",
            "summary": "No readable text was found in the document.",
            "action_items": [],
            "deadline": "No deadline specified",
            "confidence": 0.0,
            "analysis_failed": False,
            "partial_analysis": True,
            "error": "No readable text",
        }

    evidence_scores = calculate_evidence_score(text)

    print("\n🔬 COMPLETE DOCUMENT EVIDENCE SCORES")

    for category, score in evidence_scores.items():
        print(f"   {category}: {score:.3f}")

    chunks = create_chunks(text)

    print(
        f"\n📦 Document split into "
        f"{len(chunks)} chunk(s)"
    )

    if not chunks:
        chunks = [{
            "chunk_number": 1,
            "total_chunks": 1,
            "text": text[:CHUNK_SIZE],
        }]

    chunk_results = []

    for chunk in chunks:
        result = analyze_chunk_with_groq(
            chunk["text"],
            chunk["chunk_number"],
            chunk["total_chunks"],
        )

        chunk_results.append(result)

    successful_chunks = [
        r
        for r in chunk_results
        if r.get("success")
    ]

    print(
        f"\n✅ Successful Groq chunks: "
        f"{len(successful_chunks)}/{len(chunks)}"
    )

    synthesized = synthesize_results(
        chunk_results,
        text,
    )

    groq_category = synthesized.get(
        "category",
        "Other",
    )

    groq_confidence = float(
        synthesized.get(
            "groq_confidence",
            0.0,
        )
    )

    # ========================================================
    # ORIGINAL KMRL HYBRID MATH — DO NOT CHANGE
    # ========================================================

    evidence_for_category = evidence_scores.get(
        groq_category,
        0.40,
    )

    evidence_quality = min(
        (evidence_for_category - 0.40) / 0.60,
        1.0,
    )

    evidence_weight = (
        0.30
        + (
            0.50
            * evidence_quality
        )
    )

    groq_weight = 1 - evidence_weight

    final_confidence = (
        (
            evidence_weight
            * evidence_for_category
        )
        +
        (
            groq_weight
            * groq_confidence
        )
    )

    print(
        f"\n🔬 Evidence Score: "
        f"{evidence_for_category:.3f}"
    )

    print(
        f"🤖 Groq Confidence: "
        f"{groq_confidence:.3f}"
    )

    print(
        f"⚖️ Evidence Weight: "
        f"{evidence_weight:.3f}"
    )

    print(
        f"⚖️ Groq Weight: "
        f"{groq_weight:.3f}"
    )

    print(
        f"✅ Final Confidence: "
        f"{final_confidence:.3f}"
    )

    best_category = max(
        evidence_scores,
        key=evidence_scores.get,
    )

    if (
        evidence_scores.get(
            best_category,
            0,
        )
        >
        evidence_scores.get(
            groq_category,
            0,
        )
        + 0.20
    ):
        final_category = best_category

        print(
            f"⚠️ Evidence override: "
            f"{groq_category} → "
            f"{best_category}"
        )
    else:
        final_category = groq_category

    if not successful_chunks:
        final_category = best_category

        final_confidence = evidence_scores.get(
            best_category,
            0.40,
        )

        print(
            "🛡️ Groq unavailable for all chunks. "
            "Using evidence-based fallback."
        )

    partial_analysis = (
        len(successful_chunks) < len(chunks)
    )

    return {
        "category": final_category,
        "summary": synthesized.get(
            "summary",
            "No summary provided.",
        ),
        "action_items": synthesized.get(
            "action_items",
            [],
        ),
        "deadline": synthesized.get(
            "deadline",
            "No deadline specified",
        ),
        "confidence": max(
            0.0,
            min(
                float(final_confidence),
                1.0,
            ),
        ),
        "analysis_failed": False,
        "partial_analysis": partial_analysis,
        "error": None,
    }

# ============================================================
# CONVENIENCE PDF PROCESSING
# ============================================================

def process_pdf(pdf_path):
    print(f"📄 Processing: {pdf_path}")

    pages = get_page_count(pdf_path)

    print(f"📑 Pages: {pages}")

    text = extract_text_from_pdf(pdf_path)

    print(
        f"📝 Extracted text length: "
        f"{len(text)}"
    )

    analysis = analyze_with_groq(text)

    return {
        "category": analysis["category"],
        "summary": analysis["summary"],
        "action_items": analysis["action_items"],
        "deadline": analysis["deadline"],
        "pages": pages,
        "confidence": analysis["confidence"],
        "partial_analysis": analysis.get(
            "partial_analysis",
            False,
        ),
    }