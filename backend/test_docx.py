from docx import Document

doc = Document()
doc.add_paragraph("KMRL Tender Document – Track Construction.")
doc.add_paragraph("Bidding closes on 30-Aug-2026.")
doc.save("sample.docx")
print("✅ Created sample.docx")